import streamlit as st
import io
import datetime
import pandas as pd
from openpyxl import load_workbook
from utils import (
    titulo_estilizado,
    formatar_datas_dataframe,
    obter_estatisticas_mes,
    gerar_arquivo_atualizado_bytes
)

def renderizar():
    col_v1, col_v2 = st.columns([8, 2])
    with col_v2:
        if st.button("⬅️ Voltar ao Menu", key="btn_voltar_atu"):
            st.session_state["pagina"] = "menu"
            st.rerun()

    titulo_estilizado("Atualizações Gerais")

    if st.session_state.get("wb_data") is not None:
        st.info("📁 Arquivo carregado automaticamente da memória.")
        if st.checkbox("🗑️ Descartar dados da memória e carregar novo arquivo", value=False, key="desc_op2"):
            st.session_state["wb_data"] = None
            st.session_state["fila_modificacoes"] = []
            st.success("Memória limpa com sucesso!")
            st.rerun()
    else:
        st.warning("⚠️ Nenhum arquivo de destino encontrado na memória. Faça o upload abaixo.")
        sinale_file = st.file_uploader("Selecione o arquivo do SINALE (.xlsx)", type=["xlsx"], key="upload_op2")
        if sinale_file:
            st.session_state["wb_data"] = sinale_file.getvalue()
            st.session_state["last_sinale_name"] = sinale_file.name
            st.rerun()

    if st.session_state.get("wb_data") is not None:
        wb_temp = load_workbook(io.BytesIO(st.session_state["wb_data"]), data_only=True)
        target_sheet = st.selectbox("Escolha a ABA do arquivo para trabalhar:", wb_temp.sheetnames, key="aba_op2")
        header = st.number_input("Linha do cabeçalho:", value=11, min_value=1, key="header_op2")
        df = pd.read_excel(io.BytesIO(st.session_state["wb_data"]), sheet_name=target_sheet, header=header - 1)

        st.subheader("🔍 Filtros de Visualização")
        cols_para_ver = st.multiselect("Quais campos deseja visualizar?", df.columns.tolist(), default=df.columns.tolist())
        col_filtro, val_filtro = st.columns(2)
        with col_filtro:
            filtro_col = st.selectbox("Coluna para buscar:", df.columns, key="filtro_col_op2")
        valores_existentes = sorted([str(v) for v in df[filtro_col].dropna().unique()])
        with val_filtro:
            filtro_vals = st.multiselect("Selecione o(s) valor(es) para filtrar:", valores_existentes, key="filtro_vals_op2")

        df_view = df.copy()
        if filtro_vals:
            df_view = df_view[df_view[filtro_col].astype(str).isin(filtro_vals)]
        st.metric("Total de Registros Encontrados", len(df_view))

        df_view_fmt = formatar_datas_dataframe(df_view[cols_para_ver])
        st.dataframe(df_view_fmt, use_container_width=True, hide_index=True)

        st.subheader("✏️ Seleção para Atualizar")
        if "select_all" not in st.session_state:
            st.session_state["select_all"] = False
        cols_btns = st.columns([1, 1, 4])
        with cols_btns[0]:
            if st.button("✅ Marcar Todos", key="btn_marcar_t"):
                st.session_state["select_all"] = True
                st.rerun()
        with cols_btns[1]:
            if st.button("❌ Desmarcar Todos", key="btn_desmarcar_t"):
                st.session_state["select_all"] = False
                st.rerun()

        df_for_edit = df_view.copy()
        df_for_edit.insert(0, "Atualizar?", st.session_state["select_all"])
        df_editado = st.data_editor(
            df_for_edit,
            column_config={"Atualizar?": st.column_config.CheckboxColumn()},
            use_container_width=True,
            key="editor_op2"
        )

        selecionados = df_editado[df_editado["Atualizar?"] == True]
        st.metric("Total de Registros Marcados", len(selecionados))

        if not selecionados.empty:
            col_target = st.selectbox("Selecione a coluna que deseja alterar:", df.columns, key="col_target_op2")
            if col_target.strip().upper() == "DIAS":
                st.markdown("---")
                st.subheader("📅 Cálculo Automático de Dias Úteis (Seg a Sáb / Seg a Sex)")
                c_mes, c_ano = st.columns(2)
                meses_dict = {
                    "Janeiro": 1, "Fevereiro": 2, "Março": 3, "Abril": 4, "Maio": 5, "Junho": 6,
                    "Julho": 7, "Agosto": 8, "Setembro": 9, "Outubro": 10, "Novembro": 11, "Dezembro": 12
                }
                with c_mes:
                    mes_escolhido_nome = st.selectbox("Selecione o Mês:", list(meses_dict.keys()), key="sel_mes_dias")
                    mes_num = meses_dict[mes_escolhido_nome]
                with c_ano:
                    ano_escolhido = st.number_input("Digite o Ano:", min_value=2020, max_value=2035, value=datetime.date.today().year, key="sel_ano_dias")
                stats = obter_estatisticas_mes(ano_escolhido, mes_num)
                st.info(f"**Resumo para {mes_escolhido_nome}/{ano_escolhido}:**\n* **Segunda a Sábado:** {stats['seg_sab_total']} brutos | **Úteis:** **{stats['seg_sab_uteis']}**\n* **Segunda a Sexta:** {stats['seg_sex_total']} brutos | **Úteis:** **{stats['seg_sex_uteis']}**")

            valores_antigos_str = ", ".join([str(v) for v in selecionados[col_target].dropna().unique()])
            st.info(f"📌 **Valor(es) atual(is) / antigo(s)** no campo **'{col_target}'**: **{valores_antigos_str if valores_antigos_str else 'Vazio'}**")
            novo_val = st.text_input("Digite o novo valor:", key="novo_val_op2")

            if st.button("➕ Adicionar à Fila de Modificações", key="btn_add_fila"):
                st.session_state["fila_modificacoes"].append({
                    "indices": selecionados.index.tolist(),
                    "coluna": col_target,
                    "novo_valor": novo_val,
                    "valor_antigo": valores_antigos_str,
                    "vl_busca": ", ".join(filtro_vals) if filtro_vals else "Todos",
                    "aba": target_sheet
                })
                st.success("Modificação adicionada à fila!")
                st.rerun()

        if st.session_state["fila_modificacoes"]:
            st.markdown("---")
            st.subheader("📋 Fila de Modificações Pendentes")
            df_fila_resumo = pd.DataFrame([
                {
                    "Remover?": False,
                    "ID_ITEM": i,
                    "ABA": item.get("aba", "Geral"),
                    "CAMPO": item.get("coluna", ""),
                    "NOVO VALOR": item.get("novo_valor", "")
                }
                for i, item in enumerate(st.session_state["fila_modificacoes"])
            ])
            df_fila_editado = st.data_editor(
                df_fila_resumo,
                column_config={"Remover?": st.column_config.CheckboxColumn("Remover?"), "ID_ITEM": None},
                disabled=["ABA", "CAMPO", "NOVO VALOR"],
                use_container_width=True,
                key="editor_fila"
            )
            col_f1, col_f2, col_f3 = st.columns(3)
            with col_f1:
                if st.button("🗑️ Remover Selecionados"):
                    indices = df_fila_editado[df_fila_editado["Remover?"] == True]["ID_ITEM"].tolist()
                    st.session_state["fila_modificacoes"] = [
                        item for i, item in enumerate(st.session_state["fila_modificacoes"]) if i not in indices
                    ]
                    st.rerun()
            with col_f3:
                file_bytes = gerar_arquivo_atualizado_bytes(
                    io.BytesIO(st.session_state["wb_data"]),
                    header,
                    st.session_state["fila_modificacoes"],
                    df,
                    sheet_name=target_sheet
                )
                st.download_button(
                    "📥 Baixar Arquivo Atualizado",
                    file_bytes,
                    "sinale_atualizado_final.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
